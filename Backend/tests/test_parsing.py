from docx import Document as DocxDocument

from app.parsing.clause_segmenter import segment_clauses
from app.parsing.doc_parser import extract_text


def test_extract_and_segment_docx(tmp_path):
    doc = DocxDocument()
    doc.add_paragraph("1. Definitions. In this Agreement, 'Confidential Information' means any non-public information.")
    doc.add_paragraph("2. Term. This Agreement shall remain in effect for 2 years from the Effective Date.")
    doc.add_paragraph("3. Indemnification. The Vendor shall indemnify the Client against all claims.")
    path = tmp_path / "sample.docx"
    doc.save(path)

    text = extract_text(str(path))
    assert "Confidential Information" in text

    clauses = segment_clauses(text)
    assert len(clauses) == 3
    assert [c.heading for c in clauses] == ["Definitions", "Term", "Indemnification"]
    assert [c.clause_number for c in clauses] == [1, 2, 3]


def test_extract_rejects_unsupported_extension(tmp_path):
    bad = tmp_path / "sample.xyz"
    bad.write_text("hello")
    try:
        extract_text(str(bad))
        assert False, "expected ValueError for unsupported extension"
    except ValueError as e:
        assert "Unsupported file type" in str(e)


def test_segment_falls_back_to_paragraphs_when_unnumbered():
    text = "This is a plain paragraph.\n\nThis is another one."
    clauses = segment_clauses(text)
    assert len(clauses) == 2
    assert all(c.heading is None for c in clauses)