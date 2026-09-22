from unittest.mock import patch

from docx import Document as DocxDocument
from fastapi.testclient import TestClient


class ScriptedLLMClient:
    def generate_json(self, system_prompt, user_prompt, temperature=0.2):
        if "risk analyst" in system_prompt and "Indemnification" in user_prompt:
            return {"findings": [{"issue": "Uncapped indemnity.", "recommendation": "Add a cap.",
                                   "severity": "high", "confidence": 0.9}]}
        if "risk analyst" in system_prompt or "compliance analyst" in system_prompt or "drafting reviewer" in system_prompt:
            return {"findings": []}
        if "quality reviewer" in system_prompt:
            return {"decisions": [{"index": 0, "approved": True}], "needs_requeue": False}
        if "final reviewer" in system_prompt:
            return {"overall_score": 65, "overall_risk": "high",
                     "summary": "Uncapped indemnity found.", "top_issues": ["Uncapped indemnity"]}
        return {}


def _make_docx(tmp_path):
    doc = DocxDocument()
    doc.add_paragraph("1. Term. This Agreement remains in effect for 2 years.")
    doc.add_paragraph("2. Indemnification. The Vendor shall indemnify the Client against all claims, without limit.")
    path = tmp_path / "sample.docx"
    doc.save(path)
    return path


def test_health_endpoint():
    from app.main import app
    client = TestClient(app)
    resp = client.get("/health")
    assert resp.status_code == 200
    assert resp.json() == {"status": "ok"}


def test_full_upload_review_report_flow(tmp_path, clean_tables):
    path = _make_docx(tmp_path)

    with patch("app.agents.orchestrator.get_llm_client", return_value=ScriptedLLMClient()):
        from app.main import app
        client = TestClient(app)

        with open(path, "rb") as f:
            upload_resp = client.post("/documents/", files={"file": ("sample.docx", f,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        assert upload_resp.status_code == 200
        upload_body = upload_resp.json()
        assert upload_body["clause_count"] == 2
        document_id = upload_body["document_id"]

        review_resp = client.post(f"/documents/{document_id}/review")
        assert review_resp.status_code == 200
        report = review_resp.json()
        assert report["overall_score"] == 65
        assert len(report["findings"]) == 1
        assert report["findings"][0]["issue"] == "Uncapped indemnity."

        report_resp = client.get(f"/documents/{document_id}/report")
        assert report_resp.status_code == 200
        assert report_resp.json() == report


def test_review_nonexistent_document_returns_404(clean_tables):
    from app.main import app
    client = TestClient(app)
    resp = client.post("/documents/does-not-exist/review")
    assert resp.status_code == 404


def test_report_before_review_returns_404(tmp_path, clean_tables):
    path = _make_docx(tmp_path)
    from app.main import app
    client = TestClient(app)

    with open(path, "rb") as f:
        upload_resp = client.post("/documents/", files={"file": ("sample.docx", f,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    document_id = upload_resp.json()["document_id"]

    resp = client.get(f"/documents/{document_id}/report")
    assert resp.status_code == 404


def test_upload_rejects_unsupported_file_type(tmp_path, clean_tables):
    bad_file = tmp_path / "sample.xyz"
    bad_file.write_text("not a real contract")

    from app.main import app
    client = TestClient(app)
    with open(bad_file, "rb") as f:
        resp = client.post("/documents/", files={"file": ("sample.xyz", f, "application/octet-stream")})
    assert resp.status_code == 400